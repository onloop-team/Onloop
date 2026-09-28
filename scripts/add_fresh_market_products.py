#!/usr/bin/env python3
from __future__ import annotations

from copy import copy
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = ROOT / "loop_300_product_catalog.xlsx"
SHEET_NAME = "Product Catalog"
CATEGORY = "Fresh Market"

PRODUCTS = [
    ("Fresh Tomatoes", "Vegetables", "1 kg", 1500, "fresh-tomatoes-1kg.png"),
    ("Red Onions", "Vegetables", "1 kg", 1600, "fresh-red-onions-1kg.jpg"),
    ("Scotch Bonnet Pepper", "Peppers", "1 kg", 6500, "scotch-bonnet-pepper-1kg.webp"),
    ("Red Bell Pepper (Tatashe)", "Peppers", "1 kg", 6500, "red-bell-pepper-tatashe-1kg.jpg"),
    ("Cayenne Pepper (Shombo)", "Peppers", "1 kg", 5500, "cayenne-pepper-shombo-1kg.jpg"),
    ("Fresh Ginger", "Vegetables", "1 kg", 3500, "fresh-ginger-1kg.jpg"),
    ("Fresh Garlic", "Vegetables", "1 kg", 4500, "fresh-garlic-1kg.jpg"),
    ("Carrots", "Vegetables", "1 kg", 1500, "carrots-1kg.jpg"),
    ("Cabbage", "Vegetables", "1 kg", 1200, "cabbage-1kg.jpg"),
    ("Cucumber", "Vegetables", "1 kg", 1500, "cucumber-1kg.jpg"),
    ("Okro", "Vegetables", "1 kg", 1300, "okro-1kg.jpeg"),
    ("Ugu (Pumpkin Leaves)", "Vegetables", "1 kg", 2000, "ugu-pumpkin-leaves-1kg.jpg"),
    ("Apples", "Fruits", "1 kg", 4500, "apples-1kg.png"),
    ("Bananas", "Fruits", "1 kg", 1800, "bananas-1kg.webp"),
    ("Oranges", "Fruits", "1 kg", 1700, "oranges-1kg.png"),
    ("Watermelon", "Fruits", "1 kg", 900, "watermelon-1kg.jpg"),
    ("Pineapple", "Fruits", "1 kg", 1400, "pineapple-1kg.jpg"),
    ("Yam", "Tubers", "1 kg", 3000, "yam-1kg.jpg"),
    ("Ripe Plantain", "Tubers", "1 kg", 2000, "ripe-plantain-1kg.webp"),
    ("Irish Potatoes", "Tubers", "1 kg", 1800, "irish-potatoes-1kg.jpg"),
    ("Sweet Potatoes", "Tubers", "1 kg", 900, "sweet-potatoes-1kg.jpg"),
    ("Packaged Chicken", "Meat & Poultry", "1 kg", 5800, "packaged-chicken-1kg.jpg"),
    ("Chicken Drumsticks", "Meat & Poultry", "1 kg", 8000, "chicken-drumsticks-1kg.jpg"),
    ("Beef", "Meat & Poultry", "1 kg", 10000, "beef-1kg.jpg"),
    ("Goat Meat", "Meat & Poultry", "1 kg", 12500, "goat-meat-1kg.jpg"),
    ("Turkey", "Meat & Poultry", "1 kg", 7500, "turkey-1kg.png"),
    ("Catfish", "Fish & Seafood", "1 kg", 5500, "catfish-1kg.png"),
]


def main() -> None:
    workbook = openpyxl.load_workbook(WORKBOOK_PATH)
    sheet = workbook[SHEET_NAME]
    headers = {str(cell.value).strip(): cell.column for cell in sheet[1] if cell.value}

    required = {
        "Product Name",
        "Brand",
        "Main Category",
        "Subcategory",
        "Size / Unit",
        "Price",
        "Product Image",
    }
    missing = required.difference(headers)
    if missing:
        raise RuntimeError(f"Missing columns: {', '.join(sorted(missing))}")

    rows_to_delete = [
        row
        for row in range(2, sheet.max_row + 1)
        if str(sheet.cell(row, headers["Main Category"]).value or "").strip() == CATEGORY
    ]
    for row in reversed(rows_to_delete):
        sheet.delete_rows(row)

    style_row = sheet.max_row
    for name, subcategory, unit, price, filename in PRODUCTS:
        target_row = sheet.max_row + 1
        for column in range(1, sheet.max_column + 1):
            source = sheet.cell(style_row, column)
            target = sheet.cell(target_row, column)
            if source.has_style:
                target._style = copy(source._style)
            target.number_format = source.number_format
            target.alignment = copy(source.alignment)
            target.protection = copy(source.protection)

        sheet.cell(target_row, headers["Product Name"], name)
        sheet.cell(target_row, headers["Brand"], "Fresh Market")
        sheet.cell(target_row, headers["Main Category"], CATEGORY)
        sheet.cell(target_row, headers["Subcategory"], subcategory)
        sheet.cell(target_row, headers["Size / Unit"], unit)
        sheet.cell(target_row, headers["Price"], price)
        sheet.cell(
            target_row,
            headers["Product Image"],
            f"images/catalog-products/{filename}",
        )

    workbook.save(WORKBOOK_PATH)
    print(f"Added {len(PRODUCTS)} {CATEGORY} products to {WORKBOOK_PATH.name}")


if __name__ == "__main__":
    main()
